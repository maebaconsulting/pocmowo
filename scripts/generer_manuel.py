# -*- coding: utf-8 -*-
"""Génère le manuel d'utilisation CAMPOST Core Banking au format Word (.docx).

Style sobre et professionnel : bleu foncé, gris, noir, blanc. Captures d'écran
réelles intégrées depuis docs/img/. Exécuter depuis la racine du projet :

    python3 scripts/generer_manuel.py
"""
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, "docs", "img")
LOGO = os.path.join(ROOT, "public", "campost-logo.png")
OUT = os.path.join(ROOT, "docs", "Manuel_CAMPOST_Core_Banking.docx")

BLEU = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x6B, 0x71, 0x78)
NOIR = RGBColor(0x11, 0x19, 0x1F)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)

NBSP = " "  # espace insécable


def fr(texte):
    """Applique l'espacement typographique français (espace insécable)."""
    return (
        texte.replace(" :", NBSP + ":")
        .replace(" ;", NBSP + ";")
        .replace(" !", NBSP + "!")
        .replace(" ?", NBSP + "?")
        .replace("« ", "«" + NBSP)
        .replace(" »", NBSP + "»")
    )


def set_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = NOIR
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_footer(doc):
    p = doc.sections[-1].footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CAMPOST Core Banking — Manuel d'utilisation   ·   page ")
    r.font.size = Pt(8)
    r.font.color.rgb = GRIS
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "6B7178")
    rpr.append(color)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    if os.path.exists(LOGO):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO, width=Cm(3.2))

    def centered(text, size, bold, color, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    centered("CAMPOST", 34, True, BLEU, 2)
    centered("Core Banking", 20, False, GRIS, 24)
    centered("Manuel d'utilisation", 18, True, NOIR, 6)
    centered("Version 1.0", 12, False, GRIS, 2)
    centered("Application desktop · Windows", 11, False, GRIS, 2)
    doc.add_page_break()


def h1(doc, numero, titre):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(fr("%d.  %s" % (numero, titre)))
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = BLEU
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1F3A5F")
    pbdr.append(bottom)
    pPr.append(pbdr)


def h2(doc, titre):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(fr(titre))
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NOIR


def body(doc, texte):
    p = doc.add_paragraph()
    p.add_run(fr(texte))


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(fr(it))


def image(doc, fichier, legende):
    chemin = os.path.join(IMG, fichier)
    if not os.path.exists(chemin):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(chemin, width=Cm(15.5))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(fr(legende))
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = GRIS
    c.paragraph_format.space_after = Pt(10)


def add_toc(doc, titres):
    h = doc.add_paragraph()
    r = h.add_run("Sommaire")
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = BLEU
    h.paragraph_format.space_after = Pt(10)
    for i, titre in enumerate(titres, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(fr("%d.  %s" % (i, titre)))
        r.font.size = Pt(11)
        r.font.color.rgb = NOIR
    doc.add_page_break()


def creds_table(doc):
    rows = [
        ("Rôle", "Identifiant", "Mot de passe"),
        ("Administrateur", "admin@campost.cm", "admin123"),
        ("Agent", "agent@campost.cm", "agent123"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    for j, val in enumerate(rows[0]):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(val)
        run.font.bold = True
        run.font.color.rgb = BLANC
        _shade(cell, "1F3A5F")
    for i in range(1, len(rows)):
        for j, val in enumerate(rows[i]):
            cell = table.rows[i].cells[j]
            cell.text = ""
            cell.paragraphs[0].add_run(val)
    doc.add_paragraph()


def build():
    doc = Document()
    set_base_style(doc)
    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)

    add_cover(doc)
    add_footer(doc)

    titres = [
        "Présentation",
        "Prérequis",
        "Installation",
        "Connexion et rôles",
        "Prise en main de l'interface",
        "Tableau de bord",
        "Clients et KYC",
        "Comptes",
        "Dépôts et retraits",
        "Historique des transactions",
        "Relevés de compte",
        "Utilisateurs",
        "Modules à venir",
        "Désinstallation",
        "Questions fréquentes",
    ]
    add_toc(doc, titres)

    h1(doc, 1, "Présentation")
    body(doc, "CAMPOST Core Banking est l'application de gestion bancaire centrale destinée aux agences de CAMPOST. Elle réunit dans un poste de travail unique la gestion de la clientèle et de sa conformité (KYC), des comptes, des dépôts et retraits, de l'historique des transactions, des relevés de compte et le pilotage de l'activité.")
    bullets(doc, [
        "Application desktop locale pour Windows : les données sont enregistrées sur le poste et l'application fonctionne sans connexion permanente.",
        "Deux rôles : l'administrateur dispose d'un accès complet, dont la gestion des utilisateurs ; l'agent gère la clientèle, les comptes et les opérations.",
        "Montants exprimés en franc CFA (XAF, zone BEAC).",
    ])

    h1(doc, 2, "Prérequis")
    bullets(doc, [
        "Windows 10 ou Windows 11, 64 bits.",
        "Composant Microsoft Edge WebView2 : présent par défaut sur les systèmes à jour, et installé automatiquement par l'installeur si nécessaire.",
        "Aucun serveur ni base de données externe : tout est local au poste.",
    ])

    h1(doc, 3, "Installation")
    body(doc, "Munissez-vous du fichier d'installation fourni par votre service informatique : CAMPOST_1.0.0_x64-setup.exe.")
    bullets(doc, [
        "Double-cliquez sur le fichier : l'assistant d'installation s'affiche en français.",
        "L'installation se fait par utilisateur, sans droits administrateur.",
        "À la fin, un raccourci CAMPOST est ajouté au menu Démarrer et au bureau.",
        "Au premier lancement, la base de données locale est créée et un jeu de données de démonstration est chargé.",
    ])
    h2(doc, "Avertissement de sécurité")
    body(doc, "L'application n'étant pas signée numériquement, Windows peut afficher « Windows a protégé votre ordinateur ». Cliquez sur « Informations complémentaires » puis « Exécuter quand même ».")
    h2(doc, "Emplacement des données")
    body(doc, "Les données sont stockées dans le dossier %APPDATA%\\com.mowobank.poc\\. Pour réinitialiser l'application à son état de démonstration, fermez-la puis supprimez le fichier mowobank_v5.db de ce dossier.")

    h1(doc, 4, "Connexion et rôles")
    body(doc, "À l'ouverture, saisissez votre identifiant et votre mot de passe, puis cliquez sur « Se connecter ». Pour une démonstration, cliquez sur l'un des comptes proposés en bas de l'écran afin de préremplir les champs.")
    creds_table(doc)
    body(doc, "Le rôle conditionne les menus visibles : la gestion des utilisateurs et des agences est réservée à l'administrateur.")
    image(doc, "01-connexion.png", "Écran de connexion avec les comptes de démonstration.")

    h1(doc, 5, "Prise en main de l'interface")
    bullets(doc, [
        "Barre latérale gauche : navigation organisée en sections (Pilotage, Opérations, Clientèle, Crédit, Administration). L'élément actif est mis en évidence.",
        "Barre supérieure : titre de la page courante, champ de recherche et notifications.",
        "Zone de travail centrale : contenu de la page, sous forme de listes, de fiches ou de tableaux.",
        "Profil et bouton de déconnexion en bas de la barre latérale.",
    ])

    h1(doc, 6, "Tableau de bord")
    body(doc, "Le tableau de bord offre une vue d'ensemble de l'activité : cartes d'indicateurs (épargne collectée, comptes actifs, clients, volume du jour), graphique d'activité des sept derniers jours, dernières transactions et actions rapides de dépôt et de retrait.")
    image(doc, "02-tableau-de-bord.png", "Tableau de bord : indicateurs, graphique et dernières transactions.")

    h1(doc, 7, "Clients et KYC")
    body(doc, "Le module Clientèle gère le répertoire des clients et leur conformité KYC, indépendamment du personnel interne.")
    bullets(doc, [
        "Recherche par nom, téléphone, pièce ou ville, et filtre par statut KYC.",
        "Création d'un client : identité complète (nom, genre, date de naissance, profession), coordonnées, pièce d'identité et adresse. Le dossier démarre au statut KYC « en attente ».",
        "Fiche client : dossier d'identité, validation ou rejet du KYC, comptes rattachés et solde cumulé.",
    ])
    body(doc, "Règle importante : l'ouverture d'un compte n'est possible que pour un client dont le KYC est vérifié.")
    image(doc, "04-fiche-client-kyc.png", "Fiche client : identité, actions KYC et comptes rattachés.")

    h1(doc, 8, "Comptes")
    bullets(doc, [
        "Liste des comptes avec recherche (numéro, titulaire, téléphone) et filtres par type et par statut.",
        "Ouverture d'un compte : titulaire, téléphone, type (épargne, courant, tontine) et dépôt initial facultatif. L'ouverture peut aussi se faire depuis la fiche client, rattachée à ce client.",
        "Fiche compte : solde, chronologie des mouvements, actions de dépôt et de retrait, gel, réactivation, clôture et édition du relevé.",
    ])
    image(doc, "05-comptes.png", "Liste des comptes avec recherche et filtres.")
    image(doc, "06-fiche-compte.png", "Fiche compte : solde, mouvements et actions.")

    h1(doc, 9, "Dépôts et retraits")
    body(doc, "Les opérations sont accessibles depuis le tableau de bord, la fiche compte ou la page Transactions.")
    bullets(doc, [
        "Sélectionnez le compte, choisissez le type d'opération, saisissez le montant et un libellé, puis validez.",
        "Le solde et les indicateurs se mettent à jour immédiatement.",
        "Contrôle : un retrait supérieur au solde disponible est refusé ; aucune opération n'est possible sur un compte gelé ou clôturé.",
    ])
    image(doc, "07-depot.png", "Saisie d'un dépôt avec contrôle du montant.")

    h1(doc, 10, "Historique des transactions")
    body(doc, "La page Transactions présente l'historique global de toutes les opérations. Des filtres par type (dépôt, retrait) et par compte sont disponibles, ainsi que les totaux des dépôts et des retraits. Chaque ligne indique le montant signé, le solde résultant et l'opérateur.")

    h1(doc, 11, "Relevés de compte")
    body(doc, "Un relevé de compte peut être édité depuis la fiche compte, bouton « Éditer le relevé de compte », ou depuis la fiche client, bouton « Relevé » sur chaque compte rattaché.")
    bullets(doc, [
        "Choisissez la période : tout l'historique, 30 derniers jours, 90 derniers jours ou année en cours.",
        "Le relevé présente l'en-tête CAMPOST, les informations du titulaire et du compte, le solde d'ouverture, le détail des mouvements, les totaux et le solde de clôture.",
        "Le bouton « Imprimer le relevé » ouvre la fenêtre d'impression du système, qui permet également l'export en PDF.",
    ])
    image(doc, "08-releve.png", "Relevé de compte prêt à être imprimé ou exporté en PDF.")

    h1(doc, 12, "Utilisateurs")
    body(doc, "Réservée à l'administrateur, la page Utilisateurs liste le personnel de l'établissement. Elle permet de créer un utilisateur (nom, courriel, rôle, mot de passe) et de suspendre ou réactiver un compte.")

    h1(doc, 13, "Modules à venir")
    body(doc, "La plateforme prévoit d'autres modules core banking, présentés dans le menu avec leur périmètre fonctionnel : Reporting, Caisse, Virements, Groupes et tontines, Épargne, Prêts et microcrédits, Échéances, Agences, Conformité et Paramètres.")

    h1(doc, 14, "Désinstallation")
    bullets(doc, [
        "Ouvrez Paramètres Windows, puis Applications, recherchez CAMPOST et cliquez sur Désinstaller. Le désinstalleur est aussi accessible depuis le menu Démarrer.",
        "Pour supprimer également les données locales, effacez le dossier %APPDATA%\\com.mowobank.poc\\.",
    ])

    h1(doc, 15, "Questions fréquentes")
    h2(doc, "L'application fonctionne-t-elle sans Internet ?")
    body(doc, "Oui. Toutes les données et le traitement sont locaux au poste. Une connexion n'est nécessaire qu'à l'installation, et uniquement si le composant WebView2 doit être téléchargé.")
    h2(doc, "Comment réinitialiser les données de démonstration ?")
    body(doc, "Fermez l'application et supprimez le fichier de base de données dans %APPDATA%\\com.mowobank.poc\\. Il sera recréé au prochain lancement.")
    h2(doc, "J'ai oublié un mot de passe.")
    body(doc, "Un administrateur peut recréer le compte utilisateur concerné depuis la page Utilisateurs.")
    h2(doc, "Où sont stockées mes données ?")
    body(doc, "Exclusivement sur le poste, dans le dossier %APPDATA%\\com.mowobank.poc\\. Aucune donnée n'est envoyée vers un serveur externe.")

    doc.save(OUT)
    print("Manuel généré :", OUT)


if __name__ == "__main__":
    build()
